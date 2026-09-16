import io

import docx

from app.schemas.lesson_plan_content import HomeworkContent, LessonPlanContent, TeacherScriptSection, WorksheetContent


def render_lesson_plan_docx(content: LessonPlanContent, subject_name: str, year_group_name: str) -> bytes:
    document = docx.Document()
    document.add_heading(content.title, level=0)
    document.add_paragraph(f"{subject_name} · {year_group_name}")

    def section(heading: str, body):
        document.add_heading(heading, level=2)
        if isinstance(body, list):
            if body:
                for item in body:
                    document.add_paragraph(item, style="List Bullet")
            else:
                document.add_paragraph("-")
        else:
            document.add_paragraph(body or "-")

    def script_section(heading: str, script: TeacherScriptSection | None, fallback: str):
        document.add_heading(heading, level=2)
        if script is None:
            document.add_paragraph(fallback or "-")
            return
        if script.teacher_says:
            p = document.add_paragraph()
            p.add_run("TEACHER SAYS: ").bold = True
            p.add_run(f"“{script.teacher_says}”")
        if script.do:
            p = document.add_paragraph()
            p.add_run("DO: ").bold = True
            p.add_run(script.do)
        for i, question in enumerate(script.ask):
            p = document.add_paragraph()
            p.add_run("ASK: ").bold = True
            p.add_run(question)
            if i < len(script.expected_answers):
                p2 = document.add_paragraph()
                p2.add_run("EXPECTED ANSWER: ").bold = True
                p2.add_run(script.expected_answers[i])
        if script.students_do:
            p = document.add_paragraph()
            p.add_run("STUDENTS DO: ").bold = True
            p.add_run(script.students_do)
        if script.check_understanding:
            p = document.add_paragraph()
            p.add_run("CHECK FOR UNDERSTANDING: ").bold = True
            p.add_run(script.check_understanding)
        if script.watch_out_for:
            p = document.add_paragraph()
            p.add_run("WATCH OUT FOR: ").bold = True
            p.add_run(script.watch_out_for)

    section("Overview", content.overview)
    section("Learning objectives", content.learning_objectives)
    section("Success criteria", content.success_criteria)
    section("Key vocabulary", content.key_vocabulary)
    section("Prior knowledge", content.prior_knowledge)
    section("Resources needed", content.resources_needed)
    script_section("Starter", content.starter_script, content.starter)
    script_section("Main teaching / explanation", content.teacher_explanation_script, content.teacher_explanation)
    script_section("Guided practice", content.guided_practice_script, content.guided_practice)
    section("Independent practice", content.independent_practice)
    section("Key questions", content.key_questions)
    section("Differentiation - Support", content.differentiation.support)
    section("Differentiation - Core", content.differentiation.core)
    section("Differentiation - Greater depth", content.differentiation.greater_depth)
    section("Assessment", content.assessment)
    section("Common misconceptions", content.misconceptions)
    script_section("Plenary", content.plenary_script, content.plenary)
    section("Homework", content.homework)
    section("Cross-curricular links", content.cross_curricular_links)

    document.add_heading("Timeline", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text = "Time", "Activity", "Description"
    for entry in content.timeline:
        row = table.add_row().cells
        row[0].text = f"{entry.start_minute:02d}-{entry.end_minute:02d}"
        row[1].text = entry.activity
        row[2].text = entry.description

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_worksheet_docx(worksheet: WorksheetContent, subject_name: str, year_group_name: str) -> bytes:
    document = docx.Document()
    document.add_heading(worksheet.title, level=0)
    document.add_paragraph(f"{subject_name} · {year_group_name}")
    document.add_paragraph(worksheet.instructions)

    def section(heading: str, questions: list[str]):
        document.add_heading(heading, level=2)
        if questions:
            for q in questions:
                document.add_paragraph(q, style="List Number")
        else:
            document.add_paragraph("-")

    section("Recall questions", worksheet.recall_questions)
    section("Understanding questions", worksheet.understanding_questions)
    section("Application questions", worksheet.application_questions)
    section("Challenge questions", worksheet.challenge_questions)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_homework_docx(homework: HomeworkContent, subject_name: str, year_group_name: str) -> bytes:
    document = docx.Document()
    document.add_heading(homework.title, level=0)
    document.add_paragraph(f"{subject_name} · {year_group_name} · estimated {homework.estimated_minutes} minutes")
    document.add_paragraph(homework.instructions)
    document.add_heading("Tasks", level=2)
    if homework.tasks:
        for task in homework.tasks:
            document.add_paragraph(task, style="List Number")
    else:
        document.add_paragraph("-")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
