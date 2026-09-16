import docx
from pypdf import PdfReader

from app.export.docx import render_lesson_plan_docx
from app.export.pdf import render_lesson_plan_pdf
from app.schemas.lesson_plan_content import TeacherScriptSection
from tests.test_lesson_plans import SAMPLE_CONTENT


def test_export_pdf_contains_the_lesson_title():
    pdf_bytes = render_lesson_plan_pdf(SAMPLE_CONTENT, "Mathematics", "Year 2")
    reader = PdfReader(__import__("io").BytesIO(pdf_bytes))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Fractions" in text
    assert "Mathematics" in text


def test_export_docx_contains_all_sections():
    docx_bytes = render_lesson_plan_docx(SAMPLE_CONTENT, "Mathematics", "Year 2")
    document = docx.Document(__import__("io").BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert SAMPLE_CONTENT.title in full_text
    assert SAMPLE_CONTENT.overview in full_text
    assert SAMPLE_CONTENT.homework in full_text


def _scripted_content():
    content = SAMPLE_CONTENT.model_copy(deep=True)
    content.starter_script = TeacherScriptSection(
        teacher_says="Good morning everyone, today we learn about halves.",
        ask=["What is a half?"],
        expected_answers=["Two equal parts."],
        do="Show a paper circle.",
        students_do="Fold their own circle.",
        check_understanding="Ask a pupil to explain.",
        watch_out_for="Unequal parts mistaken for halves.",
    )
    return content


def test_export_pdf_renders_teacher_script_sections():
    pdf_bytes = render_lesson_plan_pdf(_scripted_content(), "Mathematics", "Year 2")
    reader = PdfReader(__import__("io").BytesIO(pdf_bytes))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "TEACHER SAYS" in text
    assert "Good morning everyone" in text
    assert "WATCH OUT FOR" in text


def test_export_docx_renders_teacher_script_sections():
    docx_bytes = render_lesson_plan_docx(_scripted_content(), "Mathematics", "Year 2")
    document = docx.Document(__import__("io").BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "TEACHER SAYS:" in full_text
    assert "Good morning everyone" in full_text
    assert "EXPECTED ANSWER:" in full_text


def test_export_still_works_without_script_sections():
    # SAMPLE_CONTENT has no script fields -- must fall back to plain text
    # cleanly, not crash, for older/library plans.
    pdf_bytes = render_lesson_plan_pdf(SAMPLE_CONTENT, "Mathematics", "Year 2")
    assert len(pdf_bytes) > 0
    docx_bytes = render_lesson_plan_docx(SAMPLE_CONTENT, "Mathematics", "Year 2")
    assert len(docx_bytes) > 0
