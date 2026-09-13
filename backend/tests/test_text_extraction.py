import io

import docx
from pptx import Presentation
from reportlab.pdfgen import canvas

from app.planning.text_extraction import TextExtractionError, chunk_text, clean_text, extract_text


def _make_pdf_bytes(text: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, text)
    c.save()
    return buf.getvalue()


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_pptx_bytes(slide_texts: list[str]) -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for text in slide_texts:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = text
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def test_extract_text_from_pdf():
    text = extract_text(_make_pdf_bytes("Hello from a PDF"), "application/pdf")
    assert "Hello from a PDF" in text


def test_extract_text_from_docx():
    text = extract_text(
        _make_docx_bytes(["First paragraph.", "Second paragraph."]),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_extract_text_from_pptx():
    text = extract_text(
        _make_pptx_bytes(["Slide one title", "Slide two title"]),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert "Slide one title" in text
    assert "Slide two title" in text


def test_extract_text_from_plain_text():
    text = extract_text(b"just some plain text", "text/plain")
    assert text == "just some plain text"


def test_extract_text_rejects_unsupported_content_type():
    try:
        extract_text(b"binary", "application/zip")
        assert False, "expected TextExtractionError"
    except TextExtractionError:
        pass


def test_clean_text_collapses_repeated_blank_lines():
    assert clean_text("a\n\n\n\nb\n\n") == "a\n\nb"


def test_chunk_text_packs_paragraphs_up_to_the_limit():
    chunks = chunk_text("a" * 10 + "\n\n" + "b" * 10, max_chars=15)
    assert chunks == ["a" * 10, "b" * 10]


def test_chunk_text_hard_splits_a_single_oversized_paragraph():
    chunks = chunk_text("x" * 25, max_chars=10)
    assert chunks == ["x" * 10, "x" * 10, "x" * 5]
